"""The whole run on generated PDFs, with a stand-in for the vision model (no GPU needed)."""
import csv
import json

import pypdfium2 as pdfium
import pytest
from openpyxl import load_workbook

from tablescrape.export import build
from tablescrape.pipeline import extract, select_shard
from tablescrape.vlm import Reply

TABLE_PAGE = json.dumps({"tables": [
    {"label": "Table 1", "caption": "Body mass by site", "columns": ["Site", "n", "Mass (kg)", "Note"],
     "rows": [["North", "12", "81.50", "=ref"], ["South", "007", "−3.2", ""]], "footnotes": "* adults only",
     "continued": False},
    {"label": "", "caption": "", "columns": ["x"], "rows": [["1"]], "footnotes": "", "continued": True},
]})


class FakeReader:
    name = "fake"

    def __init__(self, replies):
        self.replies = replies  # page number -> reply text
        self.calls = 0
        self.page = 0

    def read(self, image, page_text=""):
        self.calls += 1
        self.page += 1
        return Reply(text=self.replies.get(self.page, '{"tables": []}'), truncated=False, seconds=0.0)

    def free_memory(self):
        pass


def make_pdf(path, pages):
    path.parent.mkdir(parents=True, exist_ok=True)
    doc = pdfium.PdfDocument.new()
    for _ in range(pages):
        doc.new_page(612, 792)
    doc.save(str(path))
    doc.close()


@pytest.fixture
def corpus(tmp_path):
    pdfs = tmp_path / "pdfs"
    make_pdf(pdfs / "smith 2019.pdf", 2)
    (pdfs / "broken.PDF").write_bytes(b"not a pdf")
    return pdfs


def test_run_and_build(corpus, tmp_path):
    out = tmp_path / "out"
    reader = FakeReader({1: TABLE_PAGE})
    stats = extract(corpus, out, reader, dpi=50, log=lambda m: None)
    assert (stats["pages_read"], stats["tables"], stats["errors"]) == (2, 2, 1)
    assert (out / "pages/smith_2019/p0001.png").exists()      # page with tables keeps its picture
    assert not (out / "pages/smith_2019/p0002.png").exists()  # page without does not

    assert build(out) == {"pdfs": 2, "pages": 2, "tables": 2, "problems": 1}

    with (out / "index.csv").open(encoding="utf-8-sig") as f:
        index = list(csv.DictReader(f))
    assert [r["table_id"] for r in index] == ["smith_2019_p001_t1", "smith_2019_p001_t2"]
    assert index[0]["label"] == "Table 1" and index[0]["pdf"] == "smith 2019.pdf"
    assert index[0]["cells_to_check"] == "all (no text layer)"  # the generated PDFs are blank pages

    with (out / "csv/smith_2019/smith_2019_p001_t1.csv").open(encoding="utf-8-sig") as f:
        rows = list(csv.reader(f))
    assert rows == [["Site", "n", "Mass (kg)", "Note"], ["North", "12", "81.50", "=ref"], ["South", "007", "−3.2", ""]]

    wb = load_workbook(out / "tables.xlsx")
    assert wb.sheetnames == ["Index", "Documents", "T001 p1 smith_2019", "T002 p1 smith_2019", "Problems"]
    ws = wb["T001 p1 smith_2019"]
    assert ws["B5"].value == "Table 1" and ws["B6"].value == "Body mass by site"
    header = [c.value for c in ws[12]]
    assert header == ["Site", "n", "Mass (kg)", "Note"]
    north, south = ws[13], ws[14]
    assert north[1].value == 12 and north[2].value == 81.5 and north[2].number_format == "0.00"
    assert north[3].value == "=ref" and north[3].data_type == "s"  # not turned into a formula
    assert south[1].value == "007" and south[2].value == "−3.2"   # leading zero and Unicode minus stay text
    assert "cannot open" in wb["Problems"]["C2"].value


def test_rerun_resumes_without_rereading(corpus, tmp_path):
    out = tmp_path / "out"
    extract(corpus, out, FakeReader({1: TABLE_PAGE}), dpi=50, log=lambda m: None)
    again = FakeReader({})
    stats = extract(corpus, out, again, dpi=50, log=lambda m: None)
    assert again.calls == 0 and stats["pages_cached"] == 2


def test_bad_reply_is_recorded_and_retried(corpus, tmp_path):
    out = tmp_path / "out"
    extract(corpus, out, FakeReader({1: "sorry, no idea"}), dpi=50, log=lambda m: None)
    saved = json.loads((out / "pages/smith_2019/p0001.json").read_text())
    assert saved["error"].startswith("unreadable reply") and saved["raw"] == "sorry, no idea"
    assert build(out)["problems"] == 2

    retry = FakeReader({1: TABLE_PAGE})
    extract(corpus, out, retry, dpi=50, retry_errors=True, log=lambda m: None)
    assert retry.calls == 1
    assert build(out)["tables"] == 2


def test_shards_split_without_overlap():
    pdfs = [f"{i}.pdf" for i in range(7)]
    parts = [select_shard(pdfs, (i, 3)) for i in range(3)]
    assert sorted(sum(parts, [])) == sorted(pdfs)


def read_documents(out):
    with (out / "documents.csv").open(encoding="utf-8-sig") as f:
        return {r["pdf"]: r for r in csv.DictReader(f)}


def test_every_pdf_is_accounted_for_even_without_tables(corpus, tmp_path):
    make_pdf(corpus / "sub" / "no_tables.pdf", 3)
    out = tmp_path / "out"
    extract(corpus, out, FakeReader({}), dpi=50, log=lambda m: None)  # the model finds nothing anywhere
    assert build(out) == {"pdfs": 3, "pages": 5, "tables": 0, "problems": 1}

    docs = read_documents(out)
    assert docs["sub/no_tables.pdf"]["status"] == "no tables found"
    assert docs["sub/no_tables.pdf"]["pages_read"] == "3"
    assert docs["broken.PDF"]["status"].startswith("cannot open")
    wb = load_workbook(out / "tables.xlsx")
    assert wb.sheetnames == ["Index", "Documents", "Problems"]
    assert wb["Index"].max_row == 1  # header only


def test_unfinished_pdf_says_so(corpus, tmp_path):
    out = tmp_path / "out"
    extract(corpus, out, FakeReader({}), dpi=50, log=lambda m: None)
    (out / "pages/smith_2019/p0002.json").unlink()  # as if the job hit its time limit before page 2
    build(out)
    assert read_documents(out)["smith 2019.pdf"]["status"].startswith("not finished: 1 of 2 pages")


def test_figures_read_as_tables_are_left_out_and_reported(corpus, tmp_path):
    out = tmp_path / "out"
    reply = json.dumps({"tables": [
        {"label": "Fig. 3", "caption": "A tentative phylogeny", "columns": ["a"], "rows": [["Ursus"]]},
        {"label": "", "caption": "Figure 2. Map", "columns": ["a"], "rows": [["x"]]},
        {"label": "Table 1", "caption": "Real", "columns": ["a"], "rows": [["1"]]},
    ]})
    extract(corpus, out, FakeReader({1: reply}), dpi=50, log=lambda m: None)
    assert build(out)["tables"] == 1
    wb = load_workbook(out / "tables.xlsx")
    notes = [c.value for c in wb["Problems"]["C"][1:]]
    assert any("Fig. 3" in n for n in notes) and any("Figure 2. Map" in n for n in notes)


def test_each_skipped_figure_is_reported_once(corpus, tmp_path):
    out = tmp_path / "out"
    reply = json.dumps({"tables": [{"label": "Fig. 3", "columns": ["a"], "rows": [["x"]]}]})
    extract(corpus, out, FakeReader({1: reply}), dpi=50, log=lambda m: None)
    build(out)
    notes = [c.value for c in load_workbook(out / "tables.xlsx")["Problems"]["C"][1:] if c.value and "Fig. 3" in c.value]
    assert len(notes) == 1
